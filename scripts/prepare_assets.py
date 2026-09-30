"""Extract existing research assets without changing the source files."""
from pathlib import Path
import json
import shutil

import fitz
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'static' / 'images'
OUT.mkdir(parents=True, exist_ok=True)
desktop = Path.home() / 'Desktop'
deck = Presentation(desktop / 'dsp_rollout.pptx')
manifest = {}
for task, indices in {'cola': [1, 5, 9], 'drawer': [12, 16, 20],
                      'sweep': [23, 27, 31], 'coffee': [34, 38, 42]}.items():
    manifest[task] = []
    for stage, index in zip(['before', 'disturbance', 'after'], indices):
        picture = deck.slides[0].shapes[index].image
        name = f'{task}-{stage}.{picture.ext}'
        (OUT / name).write_bytes(picture.blob)
        manifest[task].append(f'static/images/{name}')

for name, source in {
    'prediction': Path.home() / 'Downloads/dsp_method.pdf',
    'recovery': Path.home() / 'Downloads/recovery.pdf',
}.items():
    with fitz.open(source) as document:
        document[0].get_pixmap(matrix=fitz.Matrix(2.5, 2.5), alpha=False).save(OUT / f'{name}.png')
    pdf_dir = ROOT / 'static' / 'figures'
    pdf_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, pdf_dir / f'{name}.pdf')

(ROOT / 'asset_manifest.json').write_text(json.dumps(manifest, indent=2), encoding='utf-8')
print(json.dumps(manifest, indent=2))
